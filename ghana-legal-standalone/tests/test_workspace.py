import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock

from fastapi import HTTPException
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

from legal_app.config import Settings
from legal_app.documents import extract_units
from legal_app.intelligence import Intelligence
from legal_app.main import create_app


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name)
        self.app = create_app(self.path)
        self.client = TestClient(self.app)
        self.matter = self.client.post("/api/matters", json={"name": "Synthetic contract matter"}).json()

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()

    def upload(self, text, kind="judgment", matter_id="", filename="authority.txt", **extra):
        return self.client.post("/api/documents", data={
            "kind": kind, "matter_id": matter_id, "permission_confirmed": "true", **extra,
        }, files={"file": (filename, text.encode(), "text/plain")})

    def test_upload_search_and_original_download(self):
        document = self.upload("Synthetic test authority. Contract termination requires notice.").json()
        results = self.client.post("/api/search", json={"query": "contract notice"}).json()
        self.assertEqual(results[0]["document_id"], document["id"])
        self.assertEqual(results[0]["locator"], "Paragraph 1")
        self.assertEqual(self.client.get(f"/api/passages/{results[0]['chunk_id']}").json(), results[0])
        self.assertIn(b"Synthetic test authority", self.client.get(f"/api/documents/{document['id']}/download").content)

    def test_private_matters_never_enter_library_or_other_matters(self):
        other = self.client.post("/api/matters", json={"name": "Other matter"}).json()
        self.upload("Confidential breach evidence", "matter", self.matter["id"])
        self.upload("Public breach authority")
        self.upload("Other private breach evidence", "matter", other["id"])
        library = self.client.post("/api/search", json={"query": "breach"}).json()
        self.assertTrue(all(source["kind"] != "matter" for source in library))
        matter = self.client.post("/api/search", json={"query": "breach", "scope": "matter", "matter_id": self.matter["id"]}).json()
        self.assertEqual(len(matter), 1)
        self.assertIn("Confidential", matter[0]["text"])
        both = self.client.post("/api/search", json={"query": "breach", "scope": "both", "matter_id": self.matter["id"]}).json()
        self.assertEqual(len(both), 2)
        self.assertFalse(any("Other private" in source["text"] for source in both))
        self.assertEqual(self.client.post("/api/search", json={"query": "breach", "scope": "both"}).status_code, 404)

    def test_permission_type_and_scope_validation(self):
        self.assertEqual(self.upload("Case", "matter").status_code, 404)
        self.assertEqual(self.upload("Case", matter_id=self.matter["id"]).status_code, 422)
        self.assertEqual(self.upload("Case", permission_confirmed="false").status_code, 422)
        self.assertEqual(self.upload("", filename="empty.txt").status_code, 422)
        self.assertEqual(self.upload("text", filename="bad.exe").status_code, 422)
        self.assertEqual(self.upload("source", source_url="javascript:alert(1)").status_code, 422)
        self.assertEqual(self.client.post("/api/matters", json={"name": "   "}).status_code, 422)

    def test_duplicate_upload_does_not_inflate_dashboard(self):
        first = self.upload("Same authority").json()
        second = self.upload("Same authority").json()
        self.assertEqual(first["id"], second["id"])
        self.assertEqual(self.client.get("/api/dashboard").json()["judgments"], 1)

    def test_persistence_on_restart(self):
        document = self.upload("Contract obligation evidence").json()
        with TestClient(create_app(self.path)) as restarted:
            self.assertEqual(restarted.get("/api/documents").json()[0]["id"], document["id"])
            self.assertEqual(len(restarted.post("/api/search", json={"query": "obligation"}).json()), 1)

    def test_search_punctuation_and_no_matches(self):
        self.upload("Contract evidence")
        self.assertEqual(self.client.post("/api/search", json={"query": '" OR ) * --'}).status_code, 200)
        self.assertEqual(self.client.post("/api/search", json={"query": "nonexistentphrase"}).json(), [])

    def test_legacy_brief_ids_and_stop_words_do_not_drown_relevant_terms(self):
        self.upload("Whether the and of 130 132 this was", filename="irrelevant.txt")
        relevant = self.upload("Notice to rectify a breach.", filename="notice.txt").json()
        results = self.client.post("/api/search", json={"query": "Whether the breach requires notice to rectify (source_ids: 130, 132)"}).json()
        self.assertEqual([item["document_id"] for item in results], [relevant["id"]])
        self.upload("Not permitted.", filename="negative.txt")
        results = self.client.post("/api/search", json={"query": "not"}).json()
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["text"], "Not permitted.")

    def test_legacy_case_analysis_is_paused_and_brief_still_requires_documents(self):
        payload = {"matter_id": self.matter["id"], "facts": "A synthetic factual dispute.", "issues": ["Contract breach"]}
        self.assertEqual(self.client.post("/api/case-analysis", json=payload).status_code, 409)
        self.assertEqual(self.client.post("/api/case-analysis", json={**payload, "confirmed": True}).status_code, 409)
        self.assertEqual(self.client.post("/api/case-analysis/brief", json={"matter_id": self.matter["id"]}).status_code, 422)

    def test_research_saves_sources_and_does_not_call_ai_without_evidence(self):
        intelligence = Mock()
        self.app.state.intelligence = intelligence
        self.assertEqual(self.client.post("/api/research", json={"query": "contract"}).status_code, 422)
        intelligence.report.assert_not_called()
        self.upload("Synthetic contract authority")
        sources = self.client.post("/api/search", json={"query": "contract"}).json()
        intelligence.report.return_value = {"title": "Test report", "findings": [{"text": "Test finding", "source_ids": [sources[0]["chunk_id"]]}], "sources": sources, "limitations": ["Synthetic test"]}
        result = self.client.post("/api/research", json={"query": "contract"})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(self.client.get("/api/reports").json()[0]["payload"]["sources"], sources)
        self.assertEqual(self.client.get("/api/dashboard").json()["reports"], 1)

    def test_legacy_case_analysis_cannot_bypass_the_reviewed_evidence_workflow(self):
        other = self.client.post("/api/matters", json={"name": "Other private case"}).json()
        self.upload("Synthetic public contract breach judgment")
        self.upload("Private contract breach facts", "matter", self.matter["id"])
        self.upload("Other private contract breach facts", "matter", other["id"])
        intelligence = Mock()
        intelligence.report.return_value = {"title": "Comparison", "findings": [], "limitations": []}
        self.app.state.intelligence = intelligence
        response = self.client.post("/api/case-analysis", json={"matter_id": self.matter["id"],
            "facts": "Lawyer confirmed synthetic facts.", "issues": ["contract breach"], "confirmed": True})
        self.assertEqual(response.status_code, 409)
        self.assertIn('paused', response.json()['detail'])
        intelligence.report.assert_not_called()
        self.assertEqual(self.app.state.store.reports(), [])

    def test_unknown_citation_ids_are_rejected(self):
        intelligence = Intelligence(Settings(self.path, "test-key", "", []))
        intelligence.generate = Mock(return_value={"title": "Bad", "findings": [
            {"text": "Invented", "source_ids": [999]}], "limitations": []})
        with self.assertRaises(HTTPException) as raised:
            intelligence.report("contract", [{"chunk_id": 1}])
        self.assertEqual(raised.exception.status_code, 502)

    def test_missing_api_key_fails_without_network(self):
        intelligence = Intelligence(Settings(self.path, "", "", []))
        with self.assertRaises(HTTPException) as raised:
            intelligence.brief([], False)
        self.assertEqual(raised.exception.status_code, 503)

    def test_pdf_pages_docx_tables_and_text_paragraph_locators(self):
        import fitz
        from docx import Document

        with fitz.open() as pdf:
            pdf.new_page().insert_text((70, 70), "Synthetic page one")
            pdf.new_page().insert_text((70, 70), "Synthetic page two")
            units = extract_units("sample.pdf", pdf.tobytes())
        self.assertEqual([locator for locator, _ in units], ["Page 1", "Page 2"])
        document = Document()
        document.add_paragraph("Synthetic paragraph")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Synthetic table evidence"
        stream = BytesIO()
        document.save(stream)
        units = extract_units("sample.docx", stream.getvalue())
        self.assertEqual(units[0][0], "Paragraph 1")
        self.assertEqual(units[1][0], "Table 1, row 1")
        self.assertEqual(len(extract_units("test.txt", b"First\n\nSecond")), 2)


if __name__ == "__main__":
    unittest.main()
