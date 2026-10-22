import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"


class CasePDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 9)
        self.set_text_color(30, 83, 62)
        self.cell(0, 6, "GHANA LEGAL WORKSPACE | SYNTHETIC CASE FILE", new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(190, 207, 195)
        self.line(self.l_margin, 22, self.w - self.r_margin, 22)
        self.ln(6)

    def footer(self):
        self.set_y(-14)
        self.set_font("Helvetica", size=8)
        self.set_text_color(110, 115, 110)
        self.cell(0, 6, f"Fictional training material | Page {self.page_no()} of {{nb}}", align="C")


def generate():
    data = json.loads((SAMPLES / "contract_dispute.json").read_text(encoding="utf-8"))
    pdf = CasePDF()
    pdf.alias_nb_pages()
    pdf.set_margins(18, 14, 18)
    pdf.set_auto_page_break(True, margin=20)
    document = Document()
    for section in document.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.color.rgb = RGBColor.from_string("1E533E")

    text = [f"# {data['title']}", data["subtitle"], data["notice"]]
    for index, page in enumerate(data["pages"]):
        pdf.add_page()
        if index:
            document.add_page_break()
        else:
            document.add_heading(data["title"], 0)
            document.add_paragraph(data["subtitle"])
            document.add_paragraph(data["notice"])
            pdf.set_font("Helvetica", size=8)
            pdf.set_text_color(110, 85, 55)
            pdf.multi_cell(0, 4, data["notice"], new_x="LMARGIN", new_y="NEXT")
            pdf.ln(4)
        pdf.set_font("Helvetica", "B", 15)
        pdf.set_text_color(30, 83, 62)
        pdf.multi_cell(0, 7, page["heading"], new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
        document.add_heading(page["heading"], 1)
        text.append(f"\n## {page['heading']}")
        for section in page["sections"]:
            pdf.set_font("Helvetica", "B", 10)
            pdf.set_text_color(45, 55, 48)
            pdf.multi_cell(0, 5.5, section["heading"], new_x="LMARGIN", new_y="NEXT")
            document.add_heading(section["heading"], 2)
            text.append(f"\n### {section['heading']}")
            for paragraph in section["paragraphs"]:
                pdf.set_font("Helvetica", size=9.5)
                pdf.set_text_color(45, 50, 45)
                pdf.multi_cell(0, 4.8, paragraph, new_x="LMARGIN", new_y="NEXT")
                pdf.ln(2.5)
                document.add_paragraph(paragraph)
                text.append(f"\n{paragraph}")
            pdf.ln(1)
    stem = SAMPLES / "Sample_Contract_Dispute_Ghana"
    pdf.output(str(stem.with_suffix(".pdf")))
    document.save(stem.with_suffix(".docx"))
    stem.with_suffix(".md").write_text("\n".join(text) + "\n", encoding="utf-8")
    print(f"Created PDF ({pdf.page_no()} pages), DOCX and Markdown in {SAMPLES}")


if __name__ == "__main__":
    generate()
