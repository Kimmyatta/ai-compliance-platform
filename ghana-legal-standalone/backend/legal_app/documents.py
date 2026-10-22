import re
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile


MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000


def extract_units(filename: str, content: bytes) -> list[tuple[str, str]]:
    extension = Path(filename).suffix.lower()
    units = []
    if extension == ".pdf":
        import fitz

        with fitz.open(stream=content, filetype="pdf") as document:
            if document.needs_pass:
                raise ValueError("Password-protected PDFs are not supported.")
            if len(document) > 500:
                raise ValueError("Upload a PDF with no more than 500 pages.")
            units = [(f"Page {i + 1}", page.get_text()) for i, page in enumerate(document)]
    elif extension == ".docx":
        from docx import Document

        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("The expanded DOCX exceeds the 50 MB processing limit.")
        document = Document(BytesIO(content))
        # DOCX pagination depends on the renderer; retain paragraph/table locators instead.
        units = [(f"Paragraph {i + 1}", p.text) for i, p in enumerate(document.paragraphs)]
        for i, table in enumerate(document.tables):
            for j, row in enumerate(table.rows):
                units.append((f"Table {i + 1}, row {j + 1}",
                              " | ".join(cell.text for cell in row.cells)))
    elif extension == ".txt":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Save TXT files as UTF-8 before uploading.") from exc
        units = [(f"Paragraph {i + 1}", p) for i, p in enumerate(re.split(r"\n\s*\n", text))]
    else:
        raise ValueError("Upload a PDF, DOCX, or TXT file.")
    units = [(location, text.strip()) for location, text in units if text.strip()]
    if not units:
        raise ValueError("No readable text found. Scanned documents need OCR before upload.")
    if sum(len(text) for _, text in units) > MAX_TEXT_CHARS:
        raise ValueError("The extracted document exceeds the text processing limit.")
    return units


def split_units(units: list[tuple[str, str]]) -> list[tuple[str, str]]:
    chunks = []
    for location, text in units:
        # Keep each passage in one source unit so every result has a precise locator.
        for start in range(0, len(text), 1400):
            passage = text[start:start + 1600].strip()
            if passage:
                chunks.append((location, passage))
    return chunks
